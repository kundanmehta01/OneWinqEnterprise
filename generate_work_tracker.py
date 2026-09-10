import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from datetime import date

def generate_work_tracker():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Work Tracker"
    
    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # ---------------------------------------------------------
    # Styling Definitions (Clean, No fills except Status & Priority)
    # ---------------------------------------------------------
    data_font = Font(name="Segoe UI", size=10, color="1E293B")
    data_font_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    section_font = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
    
    thin_border = Border(
        left=Side(border_style="thin", color="E2E8F0"),
        right=Side(border_style="thin", color="E2E8F0"),
        top=Side(border_style="thin", color="E2E8F0"),
        bottom=Side(border_style="thin", color="E2E8F0")
    )
    
    header_border = Border(
        left=Side(border_style="thin", color="CBD5E1"),
        right=Side(border_style="thin", color="CBD5E1"),
        top=Side(border_style="medium", color="0F172A"),
        bottom=Side(border_style="medium", color="0F172A")
    )

    section_border = Border(
        left=Side(border_style="thin", color="E2E8F0"),
        right=Side(border_style="thin", color="E2E8F0"),
        top=Side(border_style="thin", color="CBD5E1"),
        bottom=Side(border_style="medium", color="0F172A")
    )

    # Status fills & fonts
    fill_completed = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    font_completed = Font(name="Segoe UI", size=10, bold=True, color="065F46")
    
    fill_inprogress = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    font_inprogress = Font(name="Segoe UI", size=10, bold=True, color="1E40AF")
    
    fill_blocked = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_blocked = Font(name="Segoe UI", size=10, bold=True, color="991B1B")
    
    fill_onhold = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_onhold = Font(name="Segoe UI", size=10, bold=True, color="92400E")
    
    fill_notstarted = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_notstarted = Font(name="Segoe UI", size=10, bold=True, color="475569")

    # Priority fills & fonts
    fill_high = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    font_high = Font(name="Segoe UI", size=10, bold=True, color="9F1239")
    
    fill_med = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_med = Font(name="Segoe UI", size=10, bold=True, color="92400E")
    
    fill_low = PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid")
    font_low = Font(name="Segoe UI", size=10, bold=True, color="115E59")

    headers = [
        "Date",
        "Module / Feature",
        "Task / Work Description",
        "Status",
        "Priority",
        "Issues / Blockers",
        "Next Step",
        "Remarks"
    ]

    # =========================================================
    # SECTION 1: BACKEND
    # =========================================================
    ws.merge_cells("A1:H1")
    s1 = ws["A1"]
    s1.value = "BACKEND"
    s1.font = section_font
    s1.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32
    for col_idx in range(1, 9):
        ws.cell(row=1, column=col_idx).border = section_border

    # Headers for Backend
    ws.append(headers)
    ws.row_dimensions[2].height = 26
    for col_idx in range(1, 9):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.border = header_border
        if col_idx in (1, 4, 5):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Backend Completed Tasks (30 entries)
    backend_tasks = [
        (
            date(2026, 9, 1),
            "Authentication & Security",
            "Implemented Argon2id password hashing, user registration, and credential validation pipeline with failed login lockout.",
            "Completed",
            "High",
            "None",
            "Maintain session rotation & security audits",
            "Implemented in server/src/modules/auth with Argon2id + Winston logging"
        ),
        (
            date(2026, 9, 1),
            "Token Family & Session Rotation",
            "Built JWT token family rotation, multi-device refresh token management, and instant device revocation on compromise.",
            "Completed",
            "High",
            "None",
            "Monitor token expiration telemetry",
            "Implemented in auth.service.js with automatic family invalidation"
        ),
        (
            date(2026, 9, 2),
            "Password Reset & Recovery",
            "Developed cryptographically secure password reset token generation, email dispatch payload, and password update verification.",
            "Completed",
            "Medium",
            "None",
            "Connect production email SMTP credentials",
            "Implemented in auth.routes.js and auth.service.js"
        ),
        (
            date(2026, 9, 2),
            "Granular RBAC Authorization",
            "Engineered system permission registry with 48 granular permissions across 16 domain modules with role-permission check middleware.",
            "Completed",
            "High",
            "None",
            "Add custom role builder frontend integration",
            "Defined in permissions.constants.js & role.service.js"
        ),
        (
            date(2026, 9, 3),
            "Roles & Permissions Management",
            "Implemented custom role creation, role-based access control (RBAC) assignment, and system role protection rules.",
            "Completed",
            "High",
            "None",
            "Validate default system roles in staging",
            "Implemented in server/src/modules/roles and permissions"
        ),
        (
            date(2026, 9, 3),
            "Team Directory & Workforce CRUD",
            "Built paginated employee directory, automated employee ID generator (OWQ-XXX), role/department links, and soft deletion/restoration.",
            "Completed",
            "High",
            "None",
            "Review bulk import capabilities",
            "Implemented in server/src/modules/team-members"
        ),
        (
            date(2026, 9, 4),
            "Department Hierarchy & Head Assignment",
            "Engineered department CRUD operations with hierarchical parent-child relationships, department heads, and workforce count rollup.",
            "Completed",
            "Medium",
            "None",
            "Link department analytics breakdown",
            "Implemented in server/src/modules/departments"
        ),
        (
            date(2026, 9, 4),
            "Enterprise Invitations Gateway",
            "Implemented cryptographic invitation token generator, token verification endpoint, bulk invite dispatch, and onboarding assignment.",
            "Completed",
            "High",
            "None",
            "Verify invite expiration and resend flow",
            "Implemented in server/src/modules/invitations"
        ),
        (
            date(2026, 9, 5),
            "Smart NFC Card Inventory",
            "Built physical smart NFC card stock management tracking 6 materials (Matte PVC, Metal, Wood, Stainless) and status states.",
            "Completed",
            "High",
            "None",
            "Connect hardware laser UID batch importer",
            "Implemented in server/src/modules/cards/card.service.js"
        ),
        (
            date(2026, 9, 5),
            "NFC Card Binding & Lifecycle",
            "Developed one-click binding and unlinking of smart NFC cards to employee profiles with audit logging and card lost/blocked flags.",
            "Completed",
            "High",
            "None",
            "Test re-binding workflows",
            "Implemented in card.controller.js and card.service.js"
        ),
        (
            date(2026, 9, 6),
            "Contactless NFC Tap Resolver Gateway",
            "Built ultra-fast public NFC tap resolver (/api/v1/public/cards/:cardUid) with tap telemetry counter & real-time redirect to live profile.",
            "Completed",
            "High",
            "None",
            "Monitor tap telemetry performance",
            "Implemented in publicProfile.routes.js and card.service.js"
        ),
        (
            date(2026, 9, 6),
            "Profile State Isolation & Deep Diff",
            "Designed state-isolated draft vs published profile architecture with automated deep JSON diff calculator for profile changes.",
            "Completed",
            "High",
            "None",
            "Verify edge case diff formatting",
            "Implemented in employeeProfile.service.js and diff.utils.js"
        ),
        (
            date(2026, 9, 7),
            "Profile Approval Moderation Queue",
            "Built admin approval queue with side-by-side visual diff comparison, approve/reject with feedback, and auto-publishing.",
            "Completed",
            "High",
            "None",
            "Connect real-time in-app notification triggers",
            "Implemented in server/src/modules/profile-approvals"
        ),
        (
            date(2026, 9, 7),
            "Public Verified Profiles Gateway",
            "Created public verified employee profile endpoint by vanity slug (/p/:slug) with privacy filtering and SEO metadata.",
            "Completed",
            "High",
            "None",
            "Add CDN edge caching headers",
            "Implemented in server/src/modules/public-profiles"
        ),
        (
            date(2026, 9, 8),
            "Dynamic QR Code Generator Engine",
            "Implemented dynamic vector SVG and high-res Data URL QR code generator for profile sharing, vCard contact downloads, and card mapping.",
            "Completed",
            "Medium",
            "None",
            "Support customized branded QR centers",
            "Implemented in utils/qrGenerator.js and public profiles"
        ),
        (
            date(2026, 9, 8),
            "Digital Card Template Builder",
            "Engineered digital card layout template engine (Classic, Modern, Minimal, Executive, Founder) with theme color palettes and field visibility.",
            "Completed",
            "Medium",
            "None",
            "Add custom CSS theme variable injection",
            "Implemented in server/src/modules/templates"
        ),
        (
            date(2026, 9, 8),
            "8-Section Company Profile Gateway",
            "Built 8-section enterprise company profile API (Hero, About, Vision, Values, Services, Team, Gallery, Contact) for corporate branding.",
            "Completed",
            "Medium",
            "None",
            "Connect client public landing page",
            "Implemented in server/src/modules/company-profile"
        ),
        (
            date(2026, 9, 9),
            "Personalized Employee Home Dashboard",
            "Built /me/home and /me/dashboard aggregating verified profile score, KPI cards, upcoming events, and active department stream.",
            "Completed",
            "High",
            "None",
            "Add weekly engagement summary rollup",
            "Implemented in server/src/modules/user-dashboard"
        ),
        (
            date(2026, 9, 9),
            "Employee Privacy Matrix & Sessions",
            "Implemented user privacy matrix toggles (public/internal phone, email, experience) and active device session inspection/remote termination.",
            "Completed",
            "High",
            "None",
            "Review session timeout policies",
            "Implemented in server/src/modules/user-settings"
        ),
        (
            date(2026, 9, 9),
            "Colleague Networking & Connection Lifecycle",
            "Built colleague discovery feed with connection badges (none, pending_sent, pending_received, connected) and request lifecycle handlers.",
            "Completed",
            "Medium",
            "None",
            "Add mutual connections suggestion algorithm",
            "Implemented in server/src/modules/connections"
        ),
        (
            date(2026, 9, 10),
            "Enterprise Events & RSVP Ticketing",
            "Engineered department/role-scoped event feed, RSVP booking, and cryptographically secure digital ticket pass generator (OWQ-EVT-XXXXXX).",
            "Completed",
            "High",
            "None",
            "Add QR ticket scanner endpoint for event check-in",
            "Implemented in server/src/modules/events"
        ),
        (
            date(2026, 9, 10),
            "Employee Organization Explorer",
            "Developed searchable organization colleague directory for employees with department/designation filters and direct contact links.",
            "Completed",
            "Medium",
            "None",
            "Add quick connection action buttons",
            "Implemented in server/src/modules/user-directory"
        ),
        (
            date(2026, 9, 10),
            "Enterprise Helpdesk & Support Ticketing",
            "Built categorized enterprise FAQ knowledge base, support ticket creation, priority grading, and ticket status lifecycle.",
            "Completed",
            "Medium",
            "None",
            "Connect admin ticket resolution responses",
            "Implemented in server/src/modules/support"
        ),
        (
            date(2026, 9, 10),
            "Visitor Telemetry & Analytics Engine",
            "Built real-time telemetry tracking for NFC card taps, QR scans, profile views, device/browser categorization, and time-series aggregation.",
            "Completed",
            "High",
            "None",
            "Connect admin analytics charting widgets",
            "Implemented in server/src/modules/analytics"
        ),
        (
            date(2026, 9, 10),
            "Immutable Security Audit Trail",
            "Engineered automatic audit trail logging actor, target, action, diff changes, IP address, and user agent for compliance.",
            "Completed",
            "High",
            "None",
            "Add CSV / JSON audit export endpoint",
            "Implemented in server/src/modules/audit-logs"
        ),
        (
            date(2026, 9, 10),
            "In-App Notification Engine",
            "Developed event-driven notifications for profile approvals, invitations, connections, and events with unread count tracking.",
            "Completed",
            "Medium",
            "None",
            "Integrate WebSocket / SSE real-time stream",
            "Implemented in server/src/modules/notifications"
        ),
        (
            date(2026, 9, 10),
            "Media Upload & Storage Provider Integration",
            "Built multipart upload pipeline with local disk and cloud storage adapters (Cloudinary/S3) for avatars, covers, and company logos.",
            "Completed",
            "Medium",
            "None",
            "Add image optimization and thumbnail generator",
            "Implemented in server/src/modules/media"
        ),
        (
            date(2026, 9, 10),
            "Organization Security & Platform Settings",
            "Implemented global platform configuration: password strength policies, session timeouts, and approval enforcement toggles.",
            "Completed",
            "Medium",
            "None",
            "Add 2FA enforcement policy toggle",
            "Implemented in server/src/modules/settings"
        ),
        (
            date(2026, 9, 10),
            "Database Schema Architecture & Seed Pipeline",
            "Created 22 optimized Mongoose domain schemas with compound indexes, plus complete automated DB seed pipeline (Super Admin, Team, Cards).",
            "Completed",
            "High",
            "None",
            "Run seed verification in staging environment",
            "Implemented in server/src/seeds & models"
        ),
        (
            date(2026, 9, 10),
            "Security & Request Middleware Pipeline",
            "Built Helmet security headers, CORS origin whitelisting, global/auth rate limiters, request correlation IDs, and centralized AppError handler.",
            "Completed",
            "High",
            "None",
            "Monitor rate limit hits in production",
            "Implemented in server/src/middlewares & app.js"
        )
    ]

    for task in backend_tasks:
        ws.append(list(task))

    backend_end_row = len(backend_tasks) + 2  # Row 32

    for row_idx in range(3, backend_end_row + 1):
        ws.row_dimensions[row_idx].height = 24
        
        # Date
        c1 = ws.cell(row=row_idx, column=1)
        c1.font = data_font
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.border = thin_border
        c1.number_format = "DD-MM-YYYY"
        
        # Module
        c2 = ws.cell(row=row_idx, column=2)
        c2.font = data_font_bold
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c2.border = thin_border
        
        # Task Description
        c3 = ws.cell(row=row_idx, column=3)
        c3.font = data_font
        c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c3.border = thin_border
        
        # Status
        c4 = ws.cell(row=row_idx, column=4)
        c4.fill = fill_completed
        c4.font = font_completed
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c4.border = thin_border
        
        # Priority
        c5 = ws.cell(row=row_idx, column=5)
        if c5.value == "High":
            c5.fill = fill_high
            c5.font = font_high
        else:
            c5.fill = fill_med
            c5.font = font_med
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.border = thin_border
        
        # Issues / Blockers
        c6 = ws.cell(row=row_idx, column=6)
        c6.font = data_font
        c6.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c6.border = thin_border
        
        # Next Step
        c7 = ws.cell(row=row_idx, column=7)
        c7.font = data_font
        c7.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c7.border = thin_border
        
        # Remarks
        c8 = ws.cell(row=row_idx, column=8)
        c8.font = data_font
        c8.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c8.border = thin_border

    # =========================================================
    # SECTION 2: FRONTEND (Below Backend)
    # =========================================================
    # Empty separator row for visual breathing room
    sep_row = backend_end_row + 1  # Row 33
    ws.row_dimensions[sep_row].height = 14

    # Frontend Section Header (Row 34)
    frontend_header_row = sep_row + 1 # Row 34
    ws.merge_cells(f"A{frontend_header_row}:H{frontend_header_row}")
    f_cell = ws[f"A{frontend_header_row}"]
    f_cell.value = "FRONTEND"
    f_cell.font = section_font
    f_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[frontend_header_row].height = 32
    for col_idx in range(1, 9):
        ws.cell(row=frontend_header_row, column=col_idx).border = section_border

    # Frontend Column Headers (Row 35)
    frontend_cols_row = frontend_header_row + 1 # Row 35
    ws.append(headers)
    ws.row_dimensions[frontend_cols_row].height = 26
    for col_idx in range(1, 9):
        cell = ws.cell(row=frontend_cols_row, column=col_idx)
        cell.font = header_font
        cell.border = header_border
        if col_idx in (1, 4, 5):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Initial 5 ready-to-use empty rows for Frontend (Rows 36 to 40)
    for r_idx in range(frontend_cols_row + 1, frontend_cols_row + 6):
        ws.row_dimensions[r_idx].height = 24
        for c_idx in range(1, 9):
            c = ws.cell(row=r_idx, column=c_idx)
            c.font = data_font
            c.border = thin_border
            if c_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.number_format = "DD-MM-YYYY"
            elif c_idx in (4, 5):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ---------------------------------------------------------
    # Freeze Panes
    # ---------------------------------------------------------
    ws.freeze_panes = "A3"

    # ---------------------------------------------------------
    # Optimal Column Widths
    # ---------------------------------------------------------
    col_widths = {
        "A": 14,  # Date
        "B": 30,  # Module / Feature
        "C": 52,  # Task / Work Description
        "D": 16,  # Status (Colored)
        "E": 14,  # Priority (Colored)
        "F": 22,  # Issues / Blockers
        "G": 36,  # Next Step
        "H": 44   # Remarks
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # ---------------------------------------------------------
    # Data Validations (Rows 3 to 2000)
    # ---------------------------------------------------------
    # Status Dropdown (Column D)
    dv_status = DataValidation(
        type="list",
        formula1='"Not Started,In Progress,Completed,Blocked,On Hold"',
        allow_blank=True
    )
    dv_status.error = "Please select a valid status from the dropdown list."
    dv_status.errorTitle = "Invalid Status"
    dv_status.prompt = "Choose task status: Not Started, In Progress, Completed, Blocked, On Hold"
    dv_status.promptTitle = "Task Status"
    ws.add_data_validation(dv_status)
    dv_status.add("D3:D2000")

    # Priority Dropdown (Column E)
    dv_priority = DataValidation(
        type="list",
        formula1='"High,Medium,Low"',
        allow_blank=True
    )
    dv_priority.error = "Please select a valid priority (High, Medium, Low)."
    dv_priority.errorTitle = "Invalid Priority"
    dv_priority.prompt = "Choose task priority: High, Medium, Low"
    dv_priority.promptTitle = "Task Priority"
    ws.add_data_validation(dv_priority)
    dv_priority.add("E3:E2000")

    # ---------------------------------------------------------
    # Conditional Formatting (Rows 3 to 2000)
    # ---------------------------------------------------------
    # Status rules
    ws.conditional_formatting.add(
        "D3:D2000",
        CellIsRule(operator="equal", formula=['"Completed"'], stopIfTrue=True, fill=fill_completed, font=font_completed)
    )
    ws.conditional_formatting.add(
        "D3:D2000",
        CellIsRule(operator="equal", formula=['"In Progress"'], stopIfTrue=True, fill=fill_inprogress, font=font_inprogress)
    )
    ws.conditional_formatting.add(
        "D3:D2000",
        CellIsRule(operator="equal", formula=['"Blocked"'], stopIfTrue=True, fill=fill_blocked, font=font_blocked)
    )
    ws.conditional_formatting.add(
        "D3:D2000",
        CellIsRule(operator="equal", formula=['"On Hold"'], stopIfTrue=True, fill=fill_onhold, font=font_onhold)
    )
    ws.conditional_formatting.add(
        "D3:D2000",
        CellIsRule(operator="equal", formula=['"Not Started"'], stopIfTrue=True, fill=fill_notstarted, font=font_notstarted)
    )

    # Priority rules
    ws.conditional_formatting.add(
        "E3:E2000",
        CellIsRule(operator="equal", formula=['"High"'], stopIfTrue=True, fill=fill_high, font=font_high)
    )
    ws.conditional_formatting.add(
        "E3:E2000",
        CellIsRule(operator="equal", formula=['"Medium"'], stopIfTrue=True, fill=fill_med, font=font_med)
    )
    ws.conditional_formatting.add(
        "E3:E2000",
        CellIsRule(operator="equal", formula=['"Low"'], stopIfTrue=True, fill=fill_low, font=font_low)
    )

    # Save Excel file
    output_path = "d:/OneWinqEnterprise/Work_Tracker.xlsx"
    wb.save(output_path)
    print(f"Successfully generated work tracker with FRONTEND section at: {output_path}")

if __name__ == "__main__":
    generate_work_tracker()
